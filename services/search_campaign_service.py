from openpyxl import load_workbook


class SearchCampaignService:

    @staticmethod
    def health():

        return {
            "status": "SUCCESS",
            "module": "SEARCH"
        }

    @staticmethod
    def read_excel(
        file_path: str,
    ):

        workbook = load_workbook(
            file_path,
            data_only=True,
        )

        account_sheet = workbook[
            "Account & Campaign"
        ]

        campaign = {
            "account_name":
                account_sheet["C4"].value,

            "customer_id":
                account_sheet["C5"].value,

            "campaign_name":
                account_sheet["C6"].value,

            "objective":
                account_sheet["C7"].value,

            "daily_budget":
                account_sheet["C8"].value,

            "currency":
                account_sheet["C9"].value,

            "language":
                account_sheet["C10"].value,

            "locations":
                (
                    account_sheet["C11"]
                    .value
                    .split(";")
                    if account_sheet["C11"].value
                    else []
                ),

            "network":
                account_sheet["C12"].value,

            "bidding_strategy":
                account_sheet["C13"].value,

            "start_date":
                str(
                    account_sheet["C14"].value
                ),

            "end_date":
                str(
                    account_sheet["C15"].value
                ),

            "status":
                account_sheet["C16"].value,

            "label":
                account_sheet["C17"].value,
        }

        return {
            "worksheets":
                workbook.sheetnames,

            "campaign":
                campaign,
        }
